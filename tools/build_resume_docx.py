#!/usr/bin/env python3
"""Build the ATS-friendly Word resume from _data/resume_ats.yml.

Usage:
    python tools/build_resume_docx.py

Design constraints (keep these):
  * Single column, no tables, no text boxes, no images, no headers/footers.
    Applicant tracking systems mis-parse all of those.
  * Real Word heading styles so the outline is machine-readable.
  * Dates on a right tab stop rather than in a second column.
"""

from __future__ import annotations

from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "_data" / "resume_ats.yml"
OUTPUT = ROOT / "assets" / "Jin-Lee-Resume.docx"

INK = RGBColor(0x14, 0x18, 0x1F)
MUTED = RGBColor(0x4A, 0x54, 0x62)
ACCENT = RGBColor(0x1F, 0x3C, 0x88)
RULE = "C9D1DC"

BODY_FONT = "Calibri"
BODY_SIZE = Pt(10)
CONTENT_WIDTH_IN = 7.4  # 8.5in page - 0.55in margins


def _set_spacing(run, twentieths: int) -> None:
    """Apply character letter-spacing (in twentieths of a point)."""
    r_pr = run._element.get_or_add_rPr()
    spacing = r_pr.makeelement(qn("w:spacing"), {})
    spacing.set(qn("w:val"), str(twentieths))
    r_pr.append(spacing)


def _bottom_border(paragraph, color: str = RULE, size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.makeelement(qn("w:pBdr"), {})
    bottom = borders.makeelement(qn("w:bottom"), {})
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def _para(doc, text="", *, style=None, size=None, bold=False, italic=False,
          color=None, before=0, after=0, align=None):
    paragraph = doc.add_paragraph(style=style)
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if align is not None:
        paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = size or BODY_SIZE
        run.font.color.rgb = color or INK
        run.font.name = BODY_FONT
    return paragraph


def _right_tab(paragraph) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(CONTENT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT
    )


def _configure_styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.03
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for name, size in (("Heading 1", 10.5), ("Heading 2", 11), ("Heading 3", 10)):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = ACCENT if name == "Heading 1" else INK
        style.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = BODY_FONT
    bullet.font.size = BODY_SIZE
    bullet.font.color.rgb = INK
    bullet.paragraph_format.left_indent = Inches(0.2)
    bullet.paragraph_format.first_line_indent = Inches(-0.15)
    bullet.paragraph_format.space_after = Pt(2)
    bullet.paragraph_format.line_spacing = 1.03


def _section_heading(doc, text: str, first: bool = False) -> None:
    paragraph = doc.add_paragraph(style="Heading 1")
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(3 if first else 8)
    fmt.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = ACCENT
    run.font.name = BODY_FONT
    _set_spacing(run, 24)
    _bottom_border(paragraph)
    _keep_with_next(paragraph)


def _bullets(doc, items) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(" ".join(str(item).split()))
        run.font.size = BODY_SIZE
        run.font.color.rgb = INK
        run.font.name = BODY_FONT


def _role_header(doc, title: str, dates: str, before: float) -> None:
    paragraph = _para(doc, before=before, after=0)
    _right_tab(paragraph)
    title_run = paragraph.add_run(title)
    title_run.bold = True
    title_run.font.size = BODY_SIZE
    title_run.font.color.rgb = INK
    title_run.font.name = BODY_FONT
    date_run = paragraph.add_run("\t" + dates)
    date_run.font.size = Pt(9.5)
    date_run.font.color.rgb = MUTED
    date_run.font.name = BODY_FONT
    _keep_with_next(paragraph)


def build() -> None:
    data = yaml.safe_load(DATA.read_text(encoding="utf-8"))
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    _configure_styles(doc)

    # ---------- header ----------
    name = _para(doc, after=1)
    name_run = name.add_run(data["name"])
    name_run.bold = True
    name_run.font.size = Pt(21)
    name_run.font.color.rgb = INK
    name_run.font.name = BODY_FONT
    _set_spacing(name_run, 4)

    _para(doc, data["headline"], size=Pt(10.5), bold=True, color=ACCENT, after=3)
    _para(doc, "  |  ".join(data["contact"]), size=Pt(9.5), color=MUTED, after=4)

    divider = _para(doc, after=2)
    _bottom_border(divider, color="1F3C88", size=12)

    # ---------- summary ----------
    _section_heading(doc, "Professional Summary", first=True)
    _para(doc, " ".join(data["summary"].split()), after=0)

    # ---------- skills ----------
    _section_heading(doc, "Technical Skills")
    for group in data["skills"]:
        paragraph = _para(doc, after=2)
        label = paragraph.add_run(f"{group['label']}: ")
        label.bold = True
        label.font.size = BODY_SIZE
        label.font.color.rgb = INK
        label.font.name = BODY_FONT
        value = paragraph.add_run(" ".join(group["items"].split()))
        value.font.size = BODY_SIZE
        value.font.color.rgb = MUTED
        value.font.name = BODY_FONT

    # ---------- impact ----------
    _section_heading(doc, data.get("impact_heading", "Selected Impact"))
    _bullets(doc, data["impact"])

    # ---------- experience ----------
    _section_heading(doc, "Professional Experience")
    for job_index, job in enumerate(data["experience"]):
        company = doc.add_paragraph(style="Heading 2")
        company.paragraph_format.space_before = Pt(2 if job_index == 0 else 7)
        company.paragraph_format.space_after = Pt(0)
        company_run = company.add_run(job["company"])
        company_run.bold = True
        company_run.font.size = Pt(11)
        company_run.font.color.rgb = INK
        company_run.font.name = BODY_FONT
        _keep_with_next(company)

        for role_index, role in enumerate(job["roles"]):
            _role_header(doc, role["title"], role["dates"], before=0 if role_index == 0 else 4)
            if role.get("focus"):
                focus = _para(doc, role["focus"], size=Pt(9.5), italic=True,
                              color=MUTED, after=1.5)
                _keep_with_next(focus)
            _bullets(doc, role["bullets"])

    # ---------- open source ----------
    _section_heading(doc, "Selected Open Source")
    for index, project in enumerate(data["open_source"]):
        paragraph = _para(doc, before=0 if index == 0 else 4, after=0)
        name_run = paragraph.add_run(project["name"])
        name_run.bold = True
        name_run.font.size = BODY_SIZE
        name_run.font.color.rgb = INK
        name_run.font.name = BODY_FONT
        if project.get("org"):
            org_run = paragraph.add_run(f" | {project['org']}")
            org_run.font.size = BODY_SIZE
            org_run.font.color.rgb = MUTED
            org_run.font.name = BODY_FONT
        _keep_with_next(paragraph)
        description = _para(doc, " ".join(project["description"].split()), after=0)
        url_run = description.add_run(f" {project['url']}")
        url_run.font.size = Pt(9.5)
        url_run.font.color.rgb = ACCENT
        url_run.font.name = BODY_FONT

    # ---------- education ----------
    _section_heading(doc, "Education")
    education = data["education"]
    _role_header(doc, education["school"], education["dates"], before=0)
    degree = education["degree"]
    if education.get("detail"):
        degree = f"{degree} | {education['detail']}"
    _para(doc, degree, after=1)
    if education.get("honors"):
        _para(doc, f"Honors: {education['honors']}", size=Pt(9.5), color=MUTED, after=0)

    # ---------- certifications ----------
    _section_heading(doc, "Certifications")
    _para(doc, " | ".join(data["certifications"]), after=0)

    # ---------- publication (optional) ----------
    if data.get("publication"):
        _section_heading(doc, "Publication")
        _para(doc, data["publication"]["title"], after=0)

    core = doc.core_properties
    core.title = f"{data['name']} \u2014 Resume"
    core.author = data["name"]
    core.subject = data["headline"]
    core.keywords = "; ".join(group["label"] for group in data["skills"])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
